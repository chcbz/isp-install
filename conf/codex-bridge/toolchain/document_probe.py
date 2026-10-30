"""Disposable readiness probe. No business templates, network, or user data."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook

root = Path('/workspace')
doc = Document()
doc.add_paragraph('文档能力核验')
doc.save(root / 'probe.docx')
assert Document(root / 'probe.docx').paragraphs[0].text == '文档能力核验'
book = Workbook()
book.active.append(['数量', '单价', '总额'])
book.active.append([12, 18, '=A2*B2'])
book.save(root / 'probe.xlsx')
opened = load_workbook(root / 'probe.xlsx', data_only=False)
assert opened.active['A2'].value == 12 and opened.active['C2'].value == '=A2*B2'
opened.close()
print('native-documents-ready')
